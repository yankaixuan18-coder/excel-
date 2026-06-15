"""数据填充: 读取导出的数据表 xlsx, 按 (ASIN, 国家) 匹配, 供复制后自动填数。

数据表形如(第一行表头, A 列是 "ASIN；国家"):
    asin_country            inventory  sales_volume  order_quantity  ...
    B0BPM2FBGG；加拿大        295        35            34              ...

国家来源优先级(见 resolve_country):
    target.country(中文) > target.country_code(如 CA) > 从 target.name 前缀解析 > 默认 美国
"""

from __future__ import annotations

import re
from typing import Any

from .config import AppConfig, Target

# 站点代码 -> 中文名(对应数据表里 "国家" 部分)
COUNTRY_MAP = {
    "US": "美国", "USA": "美国",
    "CA": "加拿大",
    "UK": "英国", "GB": "英国",
    "FR": "法国",
    "DE": "德国",
    "IT": "意大利",
    "ES": "西班牙",
}

DEFAULT_COUNTRY = "美国"


def resolve_country(target: Target) -> str:
    if target.country:
        return target.country.strip()
    if target.country_code:
        return COUNTRY_MAP.get(target.country_code.strip().upper(), target.country_code.strip())
    # 尝试从标签名前缀解析, 如 "CA外卖包-王倩" -> CA -> 加拿大
    m = re.match(r"^([A-Za-z]{2,3})", target.name.strip())
    if m:
        code = m.group(1).upper()
        if code in COUNTRY_MAP:
            return COUNTRY_MAP[code]
    return DEFAULT_COUNTRY


def load_data_table(path: str, key_sep: str = "；") -> dict[tuple[str, str], dict[str, Any]]:
    """读取 xlsx -> {(asin, country): {字段名: 值}}。"""
    import openpyxl  # 延迟导入: 不开启填数时无需安装

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    header = next(it)
    fields = [("" if h is None else str(h).strip()) for h in header]

    index: dict[tuple[str, str], dict[str, Any]] = {}
    for row in it:
        key = row[0] if row else None
        if not key or key_sep not in str(key):
            continue
        asin, country = str(key).split(key_sep, 1)
        rec = {fields[i]: row[i] for i in range(min(len(fields), len(row))) if fields[i]}
        index[(asin.strip(), country.strip())] = rec
    return index


def lookup(index: dict, asin: str, country: str) -> dict | None:
    if not asin:
        return None
    return index.get((str(asin).strip(), str(country).strip()))


def build_row_fill(
    source_grid: list[list[dict]],
    asin_col_rel: int,
    index: dict,
    country: str,
) -> list[dict | None]:
    """对源块每一行, 取其 ASIN 去数据表匹配, 返回每行的数据记录(或 None)。"""
    out: list[dict | None] = []
    for row in source_grid:
        rec = None
        if 0 <= asin_col_rel < len(row):
            asin = (row[asin_col_rel] or {}).get("value")
            rec = lookup(index, asin, country) if asin else None
        out.append(rec)
    return out


def maybe_load(cfg: AppConfig):
    """配置了 data_table 才加载, 否则返回 None(填数功能关闭)。"""
    if not cfg.data_table:
        return None
    return load_data_table(cfg.data_table, cfg.data_key_sep)
